import re
import pandas as pd
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from comparex.config_loader import load_config


INPUT_DIR = Path("data/input")
OUTPUT_DIR = Path("data/output")


def run_reconciliation(config_name, left_file_path=None, right_file_path=None, run_id=None):

    config = load_config(config_name)

    df_left = load_dataset(config, "left", left_file_path)
    df_right = load_dataset(config, "right", right_file_path)

    df_left = normalize(df_left, config)
    df_right = normalize(df_right, config)

    df_left = map_columns(df_left, config, "left")
    df_right = map_columns(df_right, config, "right")

    df_left = trim_footer(df_left, config, "left")
    df_right = trim_footer(df_right, config, "right")

    validate_columns(df_left, df_right, config)

    df_left = apply_transforms(df_left, config, "left")
    df_right = apply_transforms(df_right, config, "right")

    df_left = cast_types(df_left, config)
    df_right = cast_types(df_right, config)

    merged = join_data(df_left, df_right, config)

    diff_df = compare_data(merged, config)

    report_df = format_report(diff_df, config)

    summary_df = build_summary(merged, config)

    output_path = build_report(report_df, summary_df, config, run_id=run_id)
    apply_highlighting(output_path, diff_df, config)

    return {
        "status": "success",
        "rows": len(diff_df),
        "output": str(output_path)
    }


# -------------------------
# file loading
# -------------------------

def load_dataset(config, side, file_path=None):

    header_row = config["datasets"][side].get("header_row", 0)

    if file_path:
        return pd.read_excel(file_path, header=header_row)

    pattern = config["datasets"][side]["file_pattern"]
    files = list(INPUT_DIR.glob(pattern))

    if not files:
        raise Exception(f"No file found matching pattern: {pattern}")

    files.sort(reverse=True)
    return pd.read_excel(files[0], header=header_row)


# -------------------------
# normalization
# -------------------------

def normalize(df, config):

    df.columns = [
        re.sub(r'[^a-z0-9]+', '', c.lower().strip())
        for c in df.columns
    ]

    if config["options"].get("trim_strings"):
        for col in df.columns:
            df[col] = df[col].astype(str).str.strip()

    return df


# -------------------------
# column mapping
# -------------------------

def normalize_col_name(col):
    return re.sub(r'[^a-z0-9]+', '', col.lower().strip())


def match_column(df_cols, pattern):

    pattern_norm = normalize_col_name(pattern)

    if "*" in pattern:
        prefix = normalize_col_name(pattern.replace("*", ""))
        matches = [c for c in df_cols if c.startswith(prefix)]

        if not matches:
            return None

        if len(matches) > 1:
            print(f"WARNING: multiple matches for '{pattern}': {matches}")

        return matches[0]

    return pattern_norm if pattern_norm in df_cols else None


def map_columns(df, config, side):

    mapping = config["column_mapping"]
    rename_dict = {}

    for canonical, col_map in mapping.items():

        source_pattern = col_map.get(side)
        if not source_pattern:
            continue

        matched_col = match_column(df.columns, source_pattern)

        if matched_col:
            rename_dict[matched_col] = canonical
        else:
            print(f"WARNING: no column match for '{source_pattern}' ({side})")

    return df.rename(columns=rename_dict)


# -------------------------
# column validator
# -------------------------

def validate_columns(df_left, df_right, config):

    expected_cols = list(config["columns"].keys())

    missing_left = [c for c in expected_cols if c not in df_left.columns]
    missing_right = [c for c in expected_cols if c not in df_right.columns]

    if missing_left:
        raise Exception(f"Missing columns in LEFT file: {missing_left}")

    if missing_right:
        raise Exception(f"Missing columns in RIGHT file: {missing_right}")


# -------------------------
# transforms
# -------------------------

def apply_transforms(df, config, side):

    for col, meta in config["columns"].items():

        transform_rules = meta.get("transform")
        if not transform_rules:
            continue

        expr = transform_rules.get(side)
        if not expr:
            continue

        df[col] = df[col].apply(
            lambda x: eval(
                expr,
                {
                    "x": x,
                    "num": pd.to_numeric(x, errors="coerce"),
                    "pd": pd,
                    "str": str
                }
            )
        )

    return df


# -------------------------
# type casting
# -------------------------

def cast_types(df, config):

    for col, meta in config["columns"].items():

        if meta["type"] == "float":
            df[col] = pd.to_numeric(df[col], errors="coerce")
        else:
            df[col] = df[col].astype(str)

    return df


# -------------------------
# tolerance
# -------------------------

def resolve_tolerance(meta, config):

    tol = meta.get("tolerance")

    if tol is None:
        return 0

    if isinstance(tol, (int, float)):
        return tol

    preset = config.get("tolerance_presets", {}).get(tol)

    if not preset:
        raise Exception(f"Tolerance preset not found: {tol}")

    return preset["abs"]


# -------------------------
# join
# -------------------------

def join_data(df_left, df_right, config):

    keys = config["join"]["keys"]
    left_name = config["datasets"]["left"]["name"]
    right_name = config["datasets"]["right"]["name"]

    return df_left.merge(
        df_right,
        on=keys,
        how="outer",
        suffixes=(f"_{left_name}", f"_{right_name}"),
        indicator=True
    )


# -------------------------
# diff engine
# -------------------------

def compare_data(df, config):

    left_name = config["datasets"]["left"]["name"]
    right_name = config["datasets"]["right"]["name"]
    opts = config["options"]
    join_keys = config["join"]["keys"]

    for col, meta in config["columns"].items():

        if meta.get("ignore") or col in join_keys:
            continue

        col_l = f"{col}_{left_name}"
        col_r = f"{col}_{right_name}"

        if col_l not in df.columns or col_r not in df.columns:
            print(f"WARNING: skipping missing column pair: {col_l}, {col_r}")
            continue

        if meta["type"] == "float":
            tol = resolve_tolerance(meta, config)
            abs_diff = abs(df[col_l] - df[col_r])
            df[f"{col}_diff"] = abs_diff
            df[f"{col}_flag"] = abs_diff > tol

        else:
            s1 = df[col_l]
            s2 = df[col_r]

            if opts.get("case_insensitive_strings"):
                s1 = s1.str.lower()
                s2 = s2.str.lower()

            df[f"{col}_diff"] = (s1 != s2)

    diff_flags = []

    for col, meta in config["columns"].items():

        if meta.get("ignore") or col in join_keys:
            continue

        if meta["type"] == "float":
            diff_flags.append(f"{col}_flag")
        else:
            diff_flags.append(f"{col}_diff")

    return df[df[diff_flags].any(axis=1)]


# -------------------------
# report builder
# -------------------------

def build_report(report_df, summary_df, config, run_id=None):

    OUTPUT_DIR.mkdir(exist_ok=True)
    file_name = config["output"]["file_name"]

    if run_id:
        output_dir = OUTPUT_DIR / run_id
        output_dir.mkdir(parents=True, exist_ok=True)
        path = output_dir / file_name
    else:
        path = OUTPUT_DIR / file_name

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        report_df.to_excel(writer, sheet_name="DETAIL", index=False)
        summary_df.to_excel(writer, sheet_name="SUMMARY", index=False)

    return path


# -------------------------
# footer trimmer
# -------------------------

def trim_footer(df, config, side):

    footer_cfg = config["datasets"][side].get("footer")
    if not footer_cfg:
        return df

    if footer_cfg["type"] == "empty_key":

        key_col = footer_cfg["column"]
        if key_col not in df.columns:
            raise Exception(f"Footer key column not found: {key_col}")

        empty_mask = df[key_col].isna() | (df[key_col] == "")

        if empty_mask.any():
            first_empty_index = empty_mask.idxmax()
            df = df.loc[:first_empty_index - 1]

    return df


# -------------------------
# report formatter
# -------------------------

def format_report(df, config):

    left_name = config["datasets"]["left"]["name"]
    right_name = config["datasets"]["right"]["name"]
    col_cfg = config["columns"]
    output_cols = list(config["join"]["keys"])

    for col, meta in col_cfg.items():

        if meta.get("ignore") or col in config["join"]["keys"]:
            continue

        label = meta.get("label", col)
        col_l = f"{col}_{left_name}"
        col_r = f"{col}_{right_name}"
        diff_col = f"{col}_diff"

        if col_l not in df.columns:
            continue

        df.rename(columns={
            col_l: f"{label}_{left_name}",
            col_r: f"{label}_{right_name}"
        }, inplace=True)

        output_cols.extend([f"{label}_{left_name}", f"{label}_{right_name}", diff_col])

    return df[output_cols]


# -------------------------
# highlighting
# -------------------------

def apply_highlighting(path, df, config):

    wb = load_workbook(path)
    ws = wb.active

    left_name = config["datasets"]["left"]["name"]
    right_name = config["datasets"]["right"]["name"]

    header_index = {cell.value: i + 1 for i, cell in enumerate(ws[1])}

    for col, meta in config["columns"].items():

        if not meta.get("highlight"):
            continue

        label = meta.get("label", col)
        diff_col = f"{col}_diff"

        if diff_col not in df.columns:
            continue

        fill = PatternFill(
            start_color=meta["color"].replace("#", ""),
            end_color=meta["color"].replace("#", ""),
            fill_type="solid"
        )

        col_l_idx = header_index.get(f"{label}_{left_name}")
        col_r_idx = header_index.get(f"{label}_{right_name}")
        diff_idx = header_index.get(diff_col)

        if not col_l_idx:
            continue

        tol = resolve_tolerance(meta, config) if meta["type"] == "float" else None

        for row_idx, val in enumerate(df[diff_col], start=2):

            highlight_flag = (val > tol) if tol is not None else bool(val)

            if highlight_flag:
                ws.cell(row=row_idx, column=col_l_idx).fill = fill
                ws.cell(row=row_idx, column=col_r_idx).fill = fill
                if diff_idx:
                    ws.cell(row=row_idx, column=diff_idx).fill = fill

    wb.save(path)


# -------------------------
# summary builder
# -------------------------

def build_summary(merged_df, config):

    key = config["join"]["keys"][0]
    left_name = config["datasets"]["left"]["name"]
    right_name = config["datasets"]["right"]["name"]

    matched_df = merged_df[merged_df["_merge"] == "both"]
    missing = merged_df[merged_df["_merge"] == "left_only"][key]
    extra = merged_df[merged_df["_merge"] == "right_only"][key]

    summary_rows = [
        {"Section": "KEY STATUS", "Metric": f"Total {left_name}", "Value": merged_df[merged_df["_merge"] != "right_only"].shape[0]},
        {"Section": "KEY STATUS", "Metric": f"Total {right_name}", "Value": merged_df[merged_df["_merge"] != "left_only"].shape[0]},
        {"Section": "KEY STATUS", "Metric": "Matched", "Value": matched_df.shape[0]},
    ]

    for val in missing:
        summary_rows.append({"Section": f"Missing in {right_name}", "Metric": key, "Value": val})

    for val in extra:
        summary_rows.append({"Section": f"Extra in {right_name}", "Metric": key, "Value": val})

    for col, meta in config["columns"].items():

        if not meta.get("summary"):
            continue

        diff_col = f"{col}_diff"
        precision = meta.get("summary_precision", 2)

        grouped = (
            matched_df[diff_col]
            .round(precision)
            .value_counts()
            .sort_index()
        )

        for diff_val, count in grouped.items():
            summary_rows.append({
                "Section": f"{meta.get('label', col)} distribution",
                "Metric": diff_val,
                "Value": count
            })

    return pd.DataFrame(summary_rows)
